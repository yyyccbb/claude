#!/usr/bin/env python3
"""
build_excel_chart.py — Build a native Excel workbook with a scatter/line
chart (X = cumulative production count, Y = actual timestamp), with NG
parts overlaid as red markers. Unlike the HTML dashboard, axis intervals
here are natively adjustable by the user in Excel (right-click axis ->
"Format Axis" -> Units).

Usage:
    python build_excel_chart.py --csv "历史记录.csv" --out-xlsx out.xlsx \
        [--time-col 时间] [--result-col 测试总结果]

GOTCHAS (do not repeat these mistakes — both caused a broken/blank axis
in earlier iterations of this exact task):
  1. The Y-values Reference range MUST start at the same row as the X-values
     range (both min_row=2, skipping the header row). If the Y range
     includes the header text cell, Excel gets an extra invalid data point
     and the whole axis scale breaks.
  2. For a ScatterChart, y_axis is a NumericAxis, not a DateAxis — even
     though the cells contain datetimes. `majorUnit` on a NumericAxis is
     interpreted in the same units as the underlying values. Excel serial
     dates are in units of *days* (1.0 = 1 day), so "30 minutes" must be
     expressed as 30/1440, NOT the literal number 30. Setting majorUnit=30
     silently produces a nonsensical multi-year axis span.
"""
import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", required=True)
    ap.add_argument("--out-xlsx", required=True)
    ap.add_argument("--time-col", default="时间")
    ap.add_argument("--result-col", default="测试总结果")
    ap.add_argument("--x-major-unit", type=float, default=100,
                     help="X轴初始刻度间隔（件），用户可在Excel里自行改")
    ap.add_argument("--y-major-minutes", type=float, default=30,
                     help="Y轴初始刻度间隔（分钟），用户可在Excel里自行改")
    args = ap.parse_args()

    df = pd.read_csv(args.csv, encoding="utf-8-sig")
    if args.time_col not in df.columns or args.result_col not in df.columns:
        print(f"[错误] 列名不匹配，实际列名: {list(df.columns)}", file=sys.stderr)
        sys.exit(1)
    df = df[[args.time_col, args.result_col]].copy()
    df.columns = ["时间", "结果"]
    df["时间"] = pd.to_datetime(df["时间"], errors="coerce")
    df = df.dropna(subset=["时间"]).sort_values("时间").reset_index(drop=True)
    df["间隔_秒"] = df["时间"].diff().dt.total_seconds()
    df["序号"] = range(1, len(df) + 1)

    wb = Workbook()
    ws = wb.active
    ws.title = "数据"

    headers = ["序号", "时间", "间隔(秒)", "结果", "NG标记时间"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="002060")
        cell.alignment = Alignment(horizontal="center")

    for i, row in df.iterrows():
        r = i + 2
        ws.cell(row=r, column=1, value=int(row["序号"]))
        tcell = ws.cell(row=r, column=2, value=row["时间"].to_pydatetime())
        tcell.number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=r, column=3, value=None if pd.isna(row["间隔_秒"]) else float(row["间隔_秒"]))
        ws.cell(row=r, column=4, value=row["结果"])
        if row["结果"] == "NG":
            ngcell = ws.cell(row=r, column=5, value=row["时间"].to_pydatetime())
            ngcell.number_format = "yyyy-mm-dd hh:mm:ss"
        for c in range(1, 6):
            ws.cell(row=r, column=c).font = Font(name="Arial")

    widths = {1: 8, 2: 20, 3: 12, 4: 8, 5: 20}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    n = len(df)
    last_row = n + 1

    chart = ScatterChart()
    chart.title = "累计生产数量-时间 折线图"
    chart.style = 2
    chart.x_axis.title = "累计生产数量（件）"
    chart.y_axis.title = "时间"
    chart.x_axis.majorUnit = args.x_major_unit
    chart.y_axis.number_format = "mm-dd hh:mm"
    # NumericAxis majorUnit is in the same unit as the cell values; Excel
    # date serials are fractions of a day, so minutes -> minutes/1440.
    chart.y_axis.majorUnit = args.y_major_minutes / 1440
    chart.height = 16
    chart.width = 34

    # X and Y ranges MUST have the same row span (min_row=2, no header row).
    x_values = Reference(ws, min_col=1, min_row=2, max_row=last_row)

    y_all = Reference(ws, min_col=2, min_row=2, max_row=last_row)
    s_all = Series(y_all, x_values, title_from_data=False, title="全部记录")
    s_all.marker = Marker(symbol="circle", size=3)
    s_all.marker.graphicalProperties.solidFill = "1F61A8"
    s_all.marker.graphicalProperties.line.noFill = True
    s_all.graphicalProperties.line.width = 9525
    s_all.graphicalProperties.line.solidFill = "1F61A8"
    s_all.smooth = False
    chart.series.append(s_all)

    y_ng = Reference(ws, min_col=5, min_row=2, max_row=last_row)
    s_ng = Series(y_ng, x_values, title_from_data=False, title="NG")
    s_ng.marker = Marker(symbol="circle", size=6)
    s_ng.marker.graphicalProperties.solidFill = "C00000"
    s_ng.marker.graphicalProperties.line.solidFill = "C00000"
    s_ng.graphicalProperties.line.noFill = True
    s_ng.smooth = False
    chart.series.append(s_ng)

    ws.add_chart(chart, "H2")
    wb.save(args.out_xlsx)
    print(f"[完成] 已写出 {args.out_xlsx}（{n} 条记录）")


if __name__ == "__main__":
    main()
